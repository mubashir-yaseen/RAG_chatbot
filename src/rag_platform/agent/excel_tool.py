"""Excel file generation tool.

Generates a real .xlsx workbook from structured tabular data and saves it under a
controlled output directory. This tool does NOT perform retrieval; it only turns
already-structured data (filename, sheet_name, rows) into a workbook. Filenames are
never derived directly from user-supplied paths: on-disk files are always named by a
generated UUID token, and the user-facing filename is carried separately as metadata.
"""

import re
import uuid
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from rag_platform.agent.tools import ToolResult
from rag_platform.config import get_settings
from rag_platform.logging_config import get_logger

logger = get_logger(__name__)

# On-disk files are always named "<32 hex chars>.xlsx"; this is also the shape
# validated by resolve_generated_file_path() to prevent path traversal.
_FILE_ID_RE = re.compile(r"^[a-f0-9]{32}$")

_MAX_ROWS = 20_000  # sane guardrail; not a configurable framework knob


def sanitize_excel_filename(name: Optional[str], default: str = "generated_file") -> str:
    """Strip unsafe characters from a user/LLM-provided display filename.

    Args:
        name: Proposed filename (may be None or empty).
        default: Fallback base name if nothing usable is provided.

    Returns:
        A safe filename ending in ``.xlsx``, capped to a reasonable length.
    """
    name = (name or "").strip()
    if not name:
        name = default
    name = re.sub(r"[^A-Za-z0-9 _\-\.]", "", name)
    name = name.strip(" .") or default
    if not name.lower().endswith(".xlsx"):
        name = f"{name}.xlsx"
    return name[:150]


def _generated_files_dir() -> Path:
    """Return (creating if needed) the controlled directory for generated files."""
    settings = get_settings()
    out_dir = Path(getattr(settings, "GENERATED_FILES_DIR", "generated_files"))
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir


def resolve_generated_file_path(file_id: str) -> Optional[Path]:
    """Resolve an opaque file_id token to a path inside the controlled output directory.

    Only accepts tokens matching the UUID-hex shape produced by create_excel_tool.
    Any other input (including path-traversal attempts) resolves to None.

    Args:
        file_id: Opaque identifier previously returned by create_excel_tool.

    Returns:
        The resolved Path if it exists inside the controlled directory, else None.
    """
    if not file_id or not _FILE_ID_RE.match(file_id):
        return None
    out_dir = _generated_files_dir().resolve()
    candidate = (out_dir / f"{file_id}.xlsx").resolve()
    if candidate.parent != out_dir:
        return None
    if not candidate.exists() or not candidate.is_file():
        return None
    return candidate


def create_excel_tool(
    filename: Optional[str] = None,
    sheet_name: Optional[str] = None,
    data: Optional[list[dict[str, Any]]] = None,
) -> ToolResult:
    """Generate a real .xlsx workbook from structured tabular data.

    Args:
        filename: Desired display filename (sanitized; ``.xlsx`` appended if missing).
        sheet_name: Worksheet name (sanitized to Excel's character restrictions).
        data: List of row dicts. Column headers are the union of keys across all rows,
            in first-seen order.

    Returns:
        ToolResult. On success, ``data["generated_file"]`` carries the metadata
        (file_id, filename, sheet_name, row_count, column_count) needed to serve the
        file for download, without exposing the real filesystem path.
    """
    if not data or not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        return ToolResult(
            tool_name="create_excel",
            output="Error: 'data' must be a non-empty list of row objects to generate an Excel file.",
            success=False,
            error="Invalid or empty data",
        )

    if len(data) > _MAX_ROWS:
        return ToolResult(
            tool_name="create_excel",
            output=f"Error: too many rows ({len(data)}); the create_excel tool supports up to {_MAX_ROWS} rows.",
            success=False,
            error="Row count exceeds limit",
        )

    safe_filename = sanitize_excel_filename(filename)
    safe_sheet = (sheet_name or "Sheet1").strip()
    safe_sheet = re.sub(r"[:\\/?*\[\]]", "", safe_sheet)[:31] or "Sheet1"

    # Union of keys across all rows, preserving first-seen order.
    columns: list[str] = []
    seen: set[str] = set()
    for row in data:
        for key in row.keys():
            key_str = str(key)
            if key_str not in seen:
                seen.add(key_str)
                columns.append(key_str)

    if not columns:
        return ToolResult(
            tool_name="create_excel",
            output="Error: row objects contained no columns to write.",
            success=False,
            error="No columns derived from data",
        )

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = safe_sheet

        header_font = Font(bold=True)
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font

        for row_idx, row in enumerate(data, start=2):
            for col_idx, col_name in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row.get(col_name))

        # Reasonable column widths from content length; no broader styling system.
        for col_idx, col_name in enumerate(columns, start=1):
            max_len = len(col_name)
            for row in data:
                val = row.get(col_name)
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 40)

        out_dir = _generated_files_dir()
        file_id = uuid.uuid4().hex
        out_path = out_dir / f"{file_id}.xlsx"
        wb.save(out_path)
    except Exception as exc:
        logger.exception("Excel generation failed")
        return ToolResult(
            tool_name="create_excel",
            output=f"Failed to generate Excel file: {exc}",
            success=False,
            error=str(exc),
        )

    logger.info(
        "Generated Excel file '%s' (%d rows, %d cols) -> file_id=%s",
        safe_filename,
        len(data),
        len(columns),
        file_id,
    )

    summary = (
        f"Created Excel file '{safe_filename}' with sheet '{safe_sheet}' "
        f"containing {len(data)} row(s) and {len(columns)} column(s)."
    )
    return ToolResult(
        tool_name="create_excel",
        output=summary,
        data={
            "generated_file": {
                "file_id": file_id,
                "filename": safe_filename,
                "sheet_name": safe_sheet,
                "row_count": len(data),
                "column_count": len(columns),
            }
        },
        success=True,
    )

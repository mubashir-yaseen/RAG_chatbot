"""Unit tests for the create_excel tool: workbook generation, MCP wiring, and agent integration."""

import json
from pathlib import Path
import pytest
from openpyxl import load_workbook

from rag_platform.agent.excel_tool import (
    create_excel_tool,
    resolve_generated_file_path,
    sanitize_excel_filename,
)
from rag_platform.agent.graph import AgentResult, run_agent
from rag_platform.agent.mcp_client import McpClient, McpServerConfig
from rag_platform.config import get_settings


@pytest.fixture(autouse=True)
def _isolated_generated_files_dir(tmp_path, monkeypatch):
    """Point GENERATED_FILES_DIR at a temp directory for every test in this file."""
    monkeypatch.setenv("GENERATED_FILES_DIR", str(tmp_path / "generated_files"))
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


SAMPLE_DATA = [
    {"category": "Current deposits", "2025": 899432634},
    {"category": "Savings deposits", "2025": 1025674118},
]


# --- EXCEL TOOL UNIT TESTS ---


@pytest.mark.unit
def test_create_excel_tool_creates_valid_workbook():
    """Verify the tool writes a real, loadable .xlsx file."""
    res = create_excel_tool(filename="BAHL_2025_deposits", sheet_name="Deposits 2025", data=SAMPLE_DATA)

    assert res.success is True
    assert res.data is not None
    generated = res.data["generated_file"]
    assert generated["filename"] == "BAHL_2025_deposits.xlsx"
    assert generated["row_count"] == 2
    assert generated["column_count"] == 2

    path = resolve_generated_file_path(generated["file_id"])
    assert path is not None
    assert path.exists()

    wb = load_workbook(path)
    assert "Deposits 2025" in wb.sheetnames


@pytest.mark.unit
def test_create_excel_tool_writes_correct_headers_and_rows():
    """Verify header row and data rows match the input exactly."""
    res = create_excel_tool(filename="test", sheet_name="Sheet1", data=SAMPLE_DATA)
    path = resolve_generated_file_path(res.data["generated_file"]["file_id"])

    wb = load_workbook(path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert headers == ["category", "2025"]

    row2 = [cell.value for cell in ws[2]]
    assert row2 == ["Current deposits", 899432634]

    row3 = [cell.value for cell in ws[3]]
    assert row3 == ["Savings deposits", 1025674118]


@pytest.mark.unit
def test_create_excel_tool_result_contains_generated_file_reference():
    """Verify the ToolResult exposes structured file metadata, not just text."""
    res = create_excel_tool(filename="report", sheet_name="Data", data=SAMPLE_DATA)
    assert "generated_file" in res.data
    gf = res.data["generated_file"]
    assert set(["file_id", "filename", "sheet_name", "row_count", "column_count"]).issubset(gf.keys())


@pytest.mark.unit
def test_create_excel_tool_rejects_empty_data():
    """Verify the tool fails cleanly (not silently) on empty/missing data."""
    res = create_excel_tool(filename="empty", sheet_name="Sheet1", data=[])
    assert res.success is False
    assert "data" in res.output.lower()


@pytest.mark.unit
def test_create_excel_tool_sanitizes_filename_and_defaults_extension():
    """Verify unsafe characters are stripped and .xlsx is appended when missing."""
    assert sanitize_excel_filename("../../etc/passwd") == "etcpasswd.xlsx"
    assert sanitize_excel_filename("My Report") == "My Report.xlsx"
    assert sanitize_excel_filename(None) == "generated_file.xlsx"


@pytest.mark.unit
def test_resolve_generated_file_path_rejects_path_traversal():
    """Verify arbitrary/malicious file_id values never resolve to a path."""
    assert resolve_generated_file_path("../../etc/passwd") is None
    assert resolve_generated_file_path("not-a-uuid") is None
    assert resolve_generated_file_path("") is None


# --- MCP CLIENT WIRING TESTS ---


@pytest.mark.unit
def test_mcp_client_discovers_create_excel_tool():
    """Verify discover_tools() registers create_excel when the excel server is configured."""
    client = McpClient(servers=[McpServerConfig(name="excel")])
    tools = client.discover_tools()
    tool_names = [t.tool_name for t in tools]
    assert "create_excel" in tool_names


@pytest.mark.unit
def test_mcp_client_call_create_excel_tool():
    """Verify McpClient.call_tool dispatches create_excel end-to-end."""
    client = McpClient(servers=[McpServerConfig(name="excel")])
    client.discover_tools()

    res = client.call_tool("create_excel", {"filename": "out", "sheet_name": "S1", "data": SAMPLE_DATA})
    assert res.success is True
    assert res.data["generated_file"]["row_count"] == 2


@pytest.mark.unit
def test_mcp_client_existing_web_search_behavior_unaffected():
    """Regression guard: adding create_excel must not break mcp_web_search dispatch."""
    client = McpClient(servers=[McpServerConfig(name="web_search")])
    client.discover_tools()
    res = client.call_tool("mcp_web_search", {"query": "latest AI news 2026"})
    assert res.success is True


# --- AGENT INTEGRATION TESTS ---


@pytest.mark.unit
def test_agent_routes_to_create_excel_with_structured_router_output():
    """Verify the agent calls create_excel when the router already supplies structured JSON."""
    client = McpClient(servers=[McpServerConfig(name="excel")])

    def mock_llm(messages, **kwargs):
        for m in messages:
            if "Classify the user query" in m.get("content", ""):
                return json.dumps(
                    {
                        "decision": "mcp",
                        "tool_name": "create_excel",
                        "tool_input": json.dumps(
                            {
                                "filename": "BAHL_2025_deposits.xlsx",
                                "sheet_name": "Deposits 2025",
                                "data": SAMPLE_DATA,
                            }
                        ),
                    }
                )
        return "I created the Excel file with the requested deposit figures."

    result = run_agent(
        query="Create an Excel file containing BAHL's 2025 deposit figures.",
        llm_call_fn=mock_llm,
        mcp_client=client,
    )

    assert isinstance(result, AgentResult)
    assert "generated_file" in result.tool_outputs
    gf = result.tool_outputs["generated_file"]
    assert gf["row_count"] == 2

    path = resolve_generated_file_path(gf["file_id"])
    assert path is not None and path.exists()


@pytest.mark.unit
def test_agent_create_excel_extraction_falls_back_to_context_grounded_llm_call():
    """Verify tool_caller_node runs a grounded extraction step when tool_input isn't JSON."""
    client = McpClient(servers=[McpServerConfig(name="excel")])
    call_count = {"n": 0}

    def mock_llm(messages, **kwargs):
        call_count["n"] += 1
        for m in messages:
            content = m.get("content", "")
            if "Classify the user query" in content:
                return json.dumps(
                    {"decision": "mcp", "tool_name": "create_excel", "tool_input": "make an excel file"}
                )
            if "Extract tabular data for an Excel export" in content:
                return json.dumps(
                    {
                        "filename": "extracted.xlsx",
                        "sheet_name": "Sheet1",
                        "data": [{"category": "Total assets", "2025": 500000}],
                    }
                )
        return "Done."

    result = run_agent(
        query="Create an excel file with the total assets figure.",
        llm_call_fn=mock_llm,
        mcp_client=client,
    )

    assert "generated_file" in result.tool_outputs
    assert result.tool_outputs["generated_file"]["row_count"] == 1


@pytest.mark.unit
def test_agent_create_excel_with_no_extractable_data_fails_gracefully():
    """Verify a request with no real data produces a clean failure, not a crash or fabricated sheet."""
    client = McpClient(servers=[McpServerConfig(name="excel")])

    def mock_llm(messages, **kwargs):
        for m in messages:
            content = m.get("content", "")
            if "Classify the user query" in content:
                return json.dumps(
                    {"decision": "mcp", "tool_name": "create_excel", "tool_input": "make an excel file"}
                )
            if "Extract tabular data for an Excel export" in content:
                return json.dumps({"data": []})
        return "I could not find any concrete data to export."

    result = run_agent(
        query="Create an excel file about nothing in particular.",
        llm_call_fn=mock_llm,
        mcp_client=client,
    )

    assert "generated_file" not in result.tool_outputs
    assert isinstance(result, AgentResult)


# --- END-TO-END: FILE DOWNLOAD ENDPOINT ---


@pytest.mark.unit
def test_download_generated_file_endpoint_serves_real_xlsx():
    """End-to-end: generate a workbook, then fetch it back through the FastAPI download route."""
    from fastapi.testclient import TestClient
    from rag_platform.api.main import app

    res = create_excel_tool(filename="BAHL_2025_deposits", sheet_name="Deposits 2025", data=SAMPLE_DATA)
    file_id = res.data["generated_file"]["file_id"]

    client = TestClient(app)
    resp = client.get(f"/api/v1/files/{file_id}", params={"filename": "BAHL_2025_deposits.xlsx"})

    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "BAHL_2025_deposits.xlsx" in resp.headers.get("content-disposition", "")
    assert len(resp.content) > 0


@pytest.mark.unit
def test_download_generated_file_endpoint_404_for_unknown_id():
    """Verify the download endpoint 404s cleanly for a nonexistent/invalid file_id."""
    from fastapi.testclient import TestClient
    from rag_platform.api.main import app

    client = TestClient(app)
    resp = client.get("/api/v1/files/" + "0" * 32)
    assert resp.status_code == 404

    resp2 = client.get("/api/v1/files/../../etc/passwd")
    assert resp2.status_code in (404, 422)
